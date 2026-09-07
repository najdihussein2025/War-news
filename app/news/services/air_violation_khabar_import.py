import json
from datetime import date, datetime
from typing import BinaryIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.news.dtos import WorkbookImportRowErrorDTO, WorkbookImportSummaryDTO
from app.news.models import AirViolation, Condition, MessageStatus, RawMessage, Village
from app.news.services.red_alert_air_violation_service import RedAlertAirViolationService
from app.sources.models import Source, SourceType
from app.sources.services.red_alert_collector import classify_condition, match_village


def read_khabar_rows(stream: BinaryIO, filename: str) -> list[dict]:
    if filename.lower().endswith('.xlsx'):
        try:
            workbook = load_workbook(stream, read_only=True, data_only=True)
        except (BadZipFile, KeyError) as exc:
            raise ValueError('The Excel workbook could not be read.') from exc
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = [str(value or '').strip().casefold() for value in next(rows, ())]
            if 'khabar' not in headers:
                raise ValueError('The workbook must contain a Khabar column.')
            return [dict(zip(headers, row)) for row in rows]
        finally:
            workbook.close()
    try:
        data = json.load(stream)
    except (ValueError, UnicodeError) as exc:
        raise ValueError('The JSON file could not be read.') from exc
    if isinstance(data, dict) and data.get('type') == 'FeatureCollection':
        data = data.get('features')
    elif isinstance(data, dict):
        data = [data]
    if not isinstance(data, list):
        raise ValueError('Provide a JSON array or GeoJSON FeatureCollection.')
    result = []
    for item in data:
        if isinstance(item, dict) and item.get('type') == 'Feature':
            item = item.get('properties')
        if not isinstance(item, dict):
            raise ValueError('Each record must contain Khabar properties.')
        result.append({str(key).strip().casefold(): value for key, value in item.items()})
    return result


class AirViolationKhabarImportService:
    def __init__(self, db: Session) -> None:
        self.db = db

    def import_file(self, stream: BinaryIO, filename: str, default_date: date) -> WorkbookImportSummaryDTO:
        rows = read_khabar_rows(stream, filename)
        villages = list(self.db.scalars(select(Village)).all())
        processed = succeeded = 0
        errors = []
        for row_number, row in enumerate(rows, start=2):
            if not any(value not in (None, '') for value in row.values()):
                continue
            processed += 1
            try:
                text = row.get('khabar')
                if not isinstance(text, str) or not text.strip():
                    raise ValueError('Khabar must contain news text.')
                text = text.strip()
                condition_id = classify_condition(text)
                if condition_id not in {35, 36, 38}:
                    raise ValueError('No supported aircraft action found in Khabar.')
                matched = match_village(text, villages)
                if matched is None:
                    raise ValueError('No village could be identified in Khabar.')
                if self.db.get(Condition, condition_id) is None:
                    raise ValueError('The aircraft action is not configured.')
                village, location = matched
                value = row.get('date')
                event_date = default_date if value in (None, '') else (
                    value.date() if isinstance(value, datetime) else
                    value if isinstance(value, date) else datetime.fromisoformat(str(value)).date()
                )
                source = self.db.scalar(select(Source).where(Source.name == 'Khabar import'))
                if source is None:
                    source = Source(type=SourceType.manual, name='Khabar import', config={}, is_active=True)
                    self.db.add(source)
                    self.db.flush()
                match = RedAlertAirViolationService._match_result(
                    text=text, condition_id=condition_id, village=village, raw_location=location,
                )
                message = RawMessage(
                    source_id=source.id, raw_text=text, raw_payload={'import': 'khabar'},
                    status=MessageStatus.routed_air_violation,
                    match_result=match.model_dump(mode='json'),
                )
                self.db.add(message)
                self.db.flush()
                self.db.add(AirViolation(
                    raw_message_id=message.id, source_id=source.id, condition_id=condition_id,
                    caza_en=village.caza_en, caza_ar=village.caza_ar,
                    event_date=event_date, event_month=event_date.strftime('%B'),
                    event_time=None, khabar=text,
                ))
                self.db.commit()
                succeeded += 1
            except (ValueError, TypeError) as exc:
                self.db.rollback()
                errors.append(WorkbookImportRowErrorDTO(row=row_number, error=str(exc)))
            except SQLAlchemyError:
                self.db.rollback()
                errors.append(WorkbookImportRowErrorDTO(row=row_number, error='The record could not be saved.'))
        return WorkbookImportSummaryDTO(
            processed=processed, succeeded=succeeded, failed=len(errors), row_errors=errors,
        )
