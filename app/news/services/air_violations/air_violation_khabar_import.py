import json
from datetime import date, datetime
from typing import BinaryIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.news.dtos import WorkbookImportRowErrorDTO, WorkbookImportSummaryDTO
from app.news.models import AirViolation, Condition, MessageStatus, RawMessage, Village
from app.news.services.air_violations.red_alert_air_violation_service import RedAlertAirViolationService
from app.news.services.air_violations.air_violation_workbook_service import AirViolationWorkbookService
from app.sources.models import Source, SourceType
from app.sources.services.red_alert_collector import classify_condition
from app.news.services.air_violations.import_source_enrichment import ImportSourceLookup, import_location_text, match_import_village, source_local_datetime, researched_import_location
from app.news.repositories.air_violation_repository import air_violation_caza_labels


def parse_event_date(value: object, default_date: date) -> date:
    if value is None or (isinstance(value, str) and not value.strip()):
        return default_date
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        for pattern in ('%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d'):
            try:
                return datetime.strptime(text, pattern).date()
            except ValueError:
                continue
    raise ValueError('Date must be an Excel date, YYYY-MM-DD, or DD/MM/YYYY.')


def read_khabar_rows(stream: BinaryIO, filename: str) -> list[dict]:
    if filename.lower().endswith('.xlsx'):
        try:
            workbook = load_workbook(stream, read_only=True, data_only=True)
        except (BadZipFile, KeyError) as exc:
            raise ValueError('The Excel workbook could not be read.') from exc
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = [str(value or '').strip().casefold() for value in next(rows, ())]
            if 'khabar' not in headers:
                raise ValueError('The workbook must contain a Khabar column.')
            return [dict(zip(headers, row)) for row in rows]
        finally:
            workbook.close()
    try:
        data = json.load(stream)
    except (ValueError, UnicodeError) as exc:
        raise ValueError('The JSON file could not be read.') from exc
    if isinstance(data, dict) and data.get('type') == 'FeatureCollection':
        data = data.get('features')
    elif isinstance(data, dict):
        data = [data]
    if not isinstance(data, list):
        raise ValueError('Provide a JSON array or GeoJSON FeatureCollection.')
    result = []
    for item in data:
        if isinstance(item, dict) and item.get('type') == 'Feature':
            item = item.get('properties')
        if not isinstance(item, dict):
            raise ValueError('Each record must contain Khabar properties.')
        result.append({str(key).strip().casefold(): value for key, value in item.items()})
    return result


class AirViolationKhabarImportService:
    def __init__(self, db: Session) -> None:
        self.db = db

    def import_file(self, stream: BinaryIO, filename: str, default_date: date) -> WorkbookImportSummaryDTO:
        rows = read_khabar_rows(stream, filename)
        villages = list(self.db.scalars(select(Village)).all())
        # Matching reads every village for every row. Keep this read-only lookup
        # outside the session so row commits/rollbacks do not expire it and cause
        # thousands of individual SELECTs on the next row.
        for village in villages:
            self.db.expunge(village)
        processed = succeeded = skipped = 0
        errors = []
        source_lookup = ImportSourceLookup(self.db)
        for row_number, row in enumerate(rows, start=2):
            if not any(value not in (None, '') for value in row.values()):
                continue
            processed += 1
            try:
                text = row.get('khabar')
                if not isinstance(text, str) or not text.strip():
                    raise ValueError('Khabar must contain news text.')
                text = text.strip()
                enrichment = source_lookup.lookup(AirViolationWorkbookService._optional(row.get('link'))).copy()
                source_text = enrichment.get('text') or ''
                condition_id = classify_condition(text) or classify_condition(source_text)
                if condition_id not in {35, 36, 38}:
                    raise ValueError('No supported aircraft action found in Khabar.')
                supplied_village = AirViolationWorkbookService._optional(row.get('village'))
                matched = match_import_village(supplied_village or text, villages, row.get('link'))
                if matched is None and not supplied_village and source_text:
                    matched = match_import_village(source_text, villages)
                if self.db.get(Condition, condition_id) is None:
                    raise ValueError('The aircraft action is not configured.')
                village, location = matched if matched else (None, None)
                optional = AirViolationWorkbookService._optional
                supplied_caza = optional(row.get('caza'))
                caza_en = optional(row.get('caza_en')) or (
                    supplied_caza if supplied_caza and not any('\u0600' <= char <= '\u06ff' for char in supplied_caza) else None
                )
                caza_ar = optional(row.get('caza_ar')) or (
                    supplied_caza if supplied_caza and any('\u0600' <= char <= '\u06ff' for char in supplied_caza) else None
                )
                if not (caza_en or caza_ar) and village is not None:
                    caza_en, caza_ar = village.caza_en, village.caza_ar
                if not (caza_en or caza_ar):
                    caza_en, caza_ar = air_violation_caza_labels(
                        text, None, None, list({(item.caza_en, item.caza_ar) for item in villages}),
                    )
                published = source_local_datetime(enrichment)
                event_date = parse_event_date(row.get('date'), published.date() if published else default_date)
                event_time = AirViolationWorkbookService._time(row.get('time'))
                if event_time is None and published:
                    event_time = published.time().replace(tzinfo=None)
                enrichment['date_source'] = 'file' if optional(row.get('date')) else 'source_post' if published else 'fallback'
                source_name = AirViolationWorkbookService._optional(row.get('source')) or 'Khabar import'
                source = self.db.scalar(select(Source).where(Source.name == source_name))
                if source is None:
                    source = Source(type=SourceType.manual, name=source_name, config={}, is_active=True)
                    self.db.add(source)
                    self.db.flush()
                # Retrying a partially successful file must not duplicate rows
                # already imported. Only compare records from this importer.
                existing = self.db.execute(select(AirViolation.id).join(
                    RawMessage, RawMessage.id == AirViolation.raw_message_id,
                ).where(
                    RawMessage.raw_payload['import'].as_string() == 'khabar',
                    AirViolation.source_id == source.id,
                    AirViolation.khabar == text,
                    # A linked post keeps its identity when a later lookup
                    # recovers its timestamp. Unlinked rows still use dates.
                    *([] if optional(row.get('link')) else [
                        AirViolation.event_date == event_date,
                        AirViolation.event_time == event_time,
                    ]),
                    AirViolation.note_1 == optional(row.get('note 1')),
                    AirViolation.note_2 == optional(row.get('note 2')),
                    AirViolation.source_link == optional(row.get('link')),
                ).limit(1)).scalar_one_or_none()
                if existing is not None:
                    skipped += 1
                    continue
                match = RedAlertAirViolationService._match_result(
                    text=text, condition_id=condition_id, village=village, raw_location=location,
                )
                match_data = match.model_dump(mode='json')
                research = researched_import_location(text, row.get('link')) if not supplied_village else None
                if research and village and village.acs_code == research['acs_code']:
                    enrichment.update(research)
                    match_data['any_village_low_confidence'] = True
                    match_data['village_matches'][0].update(village_confidence=0.8, village_match_status='matched_low_confidence', village_review_required=True)
                message = RawMessage(
                    source_id=source.id, raw_text=text,
                    raw_payload={'import': 'khabar', 'filename': filename, 'row': json.loads(json.dumps(row, default=str)),
                                 'enrichment': enrichment, 'location_text': supplied_village or import_location_text(text)},
                    source_name=enrichment.get('source_name'),
                    status=MessageStatus.routed_air_violation,
                    match_result=match_data,
                )
                self.db.add(message)
                self.db.flush()
                self.db.add(AirViolation(
                    raw_message_id=message.id, source_id=source.id, condition_id=condition_id,
                    caza_en=caza_en, caza_ar=caza_ar,
                    event_date=event_date, event_month=event_date.strftime('%B'),
                    event_time=event_time, khabar=text,
                    note_1=AirViolationWorkbookService._optional(row.get('note 1')),
                    note_2=AirViolationWorkbookService._optional(row.get('note 2')),
                    source_link=AirViolationWorkbookService._optional(row.get('link')),
                ))
                self.db.commit()
                succeeded += 1
            except (ValueError, TypeError) as exc:
                self.db.rollback()
                errors.append(WorkbookImportRowErrorDTO(row=row_number, error=str(exc)))
            except SQLAlchemyError:
                self.db.rollback()
                errors.append(WorkbookImportRowErrorDTO(row=row_number, error='The record could not be saved.'))
        return WorkbookImportSummaryDTO(
            processed=processed, succeeded=succeeded, skipped=skipped, failed=len(errors), row_errors=errors,
        )
