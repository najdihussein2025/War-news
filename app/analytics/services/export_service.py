"""Excel export of the current Analytics filters."""
from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from app.analytics.dtos.filters import Filters
from app.analytics.repositories import analytics_repository as repo
from app.analytics.services import analytics_service as service

METHODOLOGY = [
    ("Headline (counted)", "verification_status in (verified, auto_processed), not deleted, not rejected, and the "
                           "incident's action family counts in the headline."),
    ("Pending review", "needs_verification incidents plus the unclassified family (Air Activity - Needs Verification, "
                       "Unclassified / Needs Review). Excluded from the headline unless 'include pending' is on."),
    ("Casualties", "Headline totals are incidents.total_deaths / total_injuries. Male/female/children come from "
                   "incident details; unspecified = max(total - male - female - children, 0). A shared bulletin toll "
                   "sits on one incident only and is counted once; it is not attributed to a specific village."),
    ("Air violations", "Counted as CNRS breaches: warplane per mohafaza in anchored 10-minute windows; drones per caza "
                       "(6-hour windows in listed cazas, otherwise once per day); helicopters per village; spy balloons "
                       "per caza. Map levels: mohafaza = all types, caza = drones/balloons/helicopters, village = helicopters."),
    ("Geography", "Villages join to the HDX/OCHA boundaries by code (villages.acs_code = geojson code)."),
    ("History", "Rows from the pre-reset archive are frozen; live data excludes incident ids already archived."),
]


def build_workbook(db: Session, flt: Filters) -> bytes:
    meta = service.get_meta(db)
    kpis = service.get_kpis(db, flt)
    by_area = repo.rank_air_areas(db, flt, flt.level, 5000) if flt.metric == "breaches" \
        else repo.rank_incident_areas(db, flt, flt.level, 5000)
    timeline = service.get_timeline(db, flt, "day")

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    bold = Font(bold=True)
    ws.append(["Metric", "Current period", "Previous period", "Change"])
    for cell in ws[1]:
        cell.font = bold
    for label, key in (("Incidents", "incidents"), ("Deaths", "deaths"), ("Injuries", "injuries"),
                       ("Villages affected", "villages_affected"), ("Cazas affected", "cazas_affected"),
                       ("Air breaches", "air_breaches")):
        k = kpis[key]
        ws.append([label, k["current"], k["previous"], k["delta"]])
    ws.append(["Pending review (incidents)", kpis["pending_review"]])
    ws.append(["Shared-toll incidents (not attributed to a village)", kpis["unattributed_shared_toll_incidents"]])
    ws.append([])
    ws.append(["Deaths: male / female / children / unspecified", kpis["deaths"]["male"], kpis["deaths"]["female"],
               kpis["deaths"]["children"], kpis["deaths"]["unspecified"]])
    ws.append(["Injuries: male / female / children / unspecified", kpis["injuries"]["male"], kpis["injuries"]["female"],
               kpis["injuries"]["children"], kpis["injuries"]["unspecified"]])
    ws.column_dimensions["A"].width = 52

    ws = wb.create_sheet("By area")
    ws.append(["Code", "Name (EN)", "Name (AR)", "Parent", "Value", "First date", "Last date", "Dominant"])
    for cell in ws[1]:
        cell.font = bold
    for r in by_area:
        ws.append([r["code"], r["name_en"], r["name_ar"], r.get("parent_name_en"),
                   r.get("breaches") if flt.metric == "breaches" else r[
                       {"casualties": "casualties"}.get(flt.metric, flt.metric)],
                   r["first_date"], r["last_date"],
                   r.get("dominant_family") or AIR_LABEL(r.get("dominant_condition"))])

    ws = wb.create_sheet("By action")
    ws.append(["Family", "Action (EN)", "Action (AR)", "Incidents", "Deaths", "Injuries"])
    for cell in ws[1]:
        cell.font = bold
    for r in repo.condition_breakdown(db, flt):
        ws.append([r["family"], r["action_en"], r["action_ar"], int(r["incidents"]), int(r["deaths"]), int(r["injuries"])])

    ws = wb.create_sheet("Timeline")
    ws.append(["Date", "Total", "7-day average"] + [s["key"] for s in timeline["series"]])
    for cell in ws[1]:
        cell.font = bold
    for i, point in enumerate(timeline["total"]):
        ws.append([point["date"], point["value"], timeline["moving_average_7d"][i]["value"]]
                  + [s["points"][i]["value"] for s in timeline["series"]])

    ws = wb.create_sheet("Methodology")
    ws.append(["Filter / rule", "Value"])
    for cell in ws[1]:
        cell.font = bold
    params = flt.cache_params()
    for key in ("date_from", "date_to", "level", "area_code", "family", "condition_id", "include_pending", "metric"):
        ws.append([f"Filter: {key}", str(params[key])])
    for label, text_ in METHODOLOGY:
        ws.append([label, text_])
    ws.append(["Archive cut-off", str(meta["archive"]["cutoff_date"] or "no archive yet")])
    ws.append(["Data last refreshed", str(meta["last_refresh"])])
    ws.append(["Export created", datetime.now(timezone.utc).isoformat()])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 120

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def AIR_LABEL(condition_id):  # noqa: N802
    from app.analytics.constants import AIR_CONDITION_LABELS

    return AIR_CONDITION_LABELS.get(condition_id) if condition_id else None
