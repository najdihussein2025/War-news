from datetime import date, time
from types import SimpleNamespace
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient

from app.analytics.services import refresh_service
from app.api.deps import require_admin
from app.core import cache
from app.core.database import get_db
from app.main import app
from tests.analytics.support import BALLOON, DRONE, HELICOPTER, WARPLANE

D1, D2, D3 = date(2026, 10, 1), date(2026, 10, 2), date(2026, 10, 3)
Q = "date_from=2026-10-01&date_to=2026-10-03"


@pytest.fixture()
def fake_cache(monkeypatch):
    store, versions = {}, {"v": 1}
    monkeypatch.setattr(cache, "get_json", lambda key: store.get(key))
    monkeypatch.setattr(cache, "set_json", lambda key, value, ttl: store.__setitem__(key, value))
    monkeypatch.setattr(cache, "get_cache_version", lambda key: versions["v"])
    monkeypatch.setattr(cache, "increment", lambda key: versions.__setitem__("v", versions["v"] + 1) or versions["v"])
    return store, versions


@pytest.fixture()
def client(db, fake_cache):
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[require_admin] = lambda: SimpleNamespace(id=uuid4())
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.clear()


@pytest.fixture()
def data(world, db):
    sour0, sour1, saida, zahle = world.village("Sour", 0), world.village("Sour", 1), world.village("Saida"), world.village("Zahle")
    world.incident(village_id=sour0, condition="Bombs", day=D1, deaths=3, injuries=2, male_d=1, female_d=1, children_d=1)
    world.incident(village_id=sour0, condition="Artillery Shelling", day=D2, deaths=1)
    world.incident(village_id=sour1, condition="Bombs", day=D2, injuries=4)
    world.incident(village_id=saida, condition="Shooting", day=D3)
    world.incident(village_id=zahle, condition="Bombs", day=D3, status="needs_verification", deaths=7)
    world.incident(village_id=zahle, condition="Bombs", day=D3, deleted=True, deaths=50)
    world.incident(village_id=sour0, condition="Bombs", day=D1, deaths=2, casualty_status="aggregate_only")
    world.air(condition_id=WARPLANE, day=D2, at=time(10, 0), village_ids=(sour0,), commit=False)
    world.air(condition_id=DRONE, day=D2, at=time(11, 0), village_ids=(sour0,), commit=False)
    world.air(condition_id=HELICOPTER, day=D3, at=time(9, 0), village_ids=(sour1,), commit=False)
    world.air(condition_id=BALLOON, day=D3, at=time(9, 0), caza_en="Aley", commit=False)
    db.commit()
    refresh_service.full_rebuild(db)
    return {"sour0": sour0}


ENDPOINTS = ["meta", "map", "kpis", "timeline", "rankings", "peaks", "matrix", "incidents", "area/LB63", "export.xlsx"]


@pytest.mark.parametrize("path", ENDPOINTS)
def test_auth_required(path):
    app.dependency_overrides.clear()
    response = TestClient(app).get(f"/api/analytics/{path}")
    assert response.status_code == 401


@pytest.mark.parametrize("query", [
    "date_from=2026-10-05&date_to=2026-10-01", "date_from=2024-01-01&date_to=2026-10-01",
    "level=planet", "metric=money", "area_code=ZZ99", "family=nope", "condition_id=99999",
])
def test_filter_validation_rejects_bad_values(client, data, query):
    assert client.get(f"/api/analytics/map?{query}").status_code == 422


def test_meta(client, data):
    body = client.get("/api/analytics/meta").json()
    assert body["area_counts"] == {"mohafaza": 8, "caza": 26, "village": 1611}
    assert body["incident_data"]["first_date"] == "2026-10-01"
    assert body["pending_review"]["incidents"] == 1
    assert len(body["families"]) == 10


def test_map_incidents_by_level(client, data):
    body = client.get(f"/api/analytics/map?{Q}&level=caza").json()
    values = {v["code"]: v["value"] for v in body["values"]}
    assert values == {"LB63": 4, "LB61": 1}  # Sour 4 headline, Saida 1; pending/deleted excluded
    assert all(set(v) <= {"code", "value", "families"} for v in body["values"])
    withp = {v["code"]: v["value"] for v in client.get(f"/api/analytics/map?{Q}&level=caza&include_pending=true").json()["values"]}
    assert withp["LB25"] == 1


def test_map_metrics_and_village_level(client, data):
    deaths = {v["code"]: v["value"] for v in client.get(f"/api/analytics/map?{Q}&level=mohafaza&metric=deaths").json()["values"]}
    assert deaths == {"LB6": 6}
    village = client.get(f"/api/analytics/map?{Q}&level=village&metric=incidents").json()["values"]
    assert sum(v["value"] for v in village) == 5


def test_map_air_rules_reject_unsupported_levels(client, data):
    assert client.get(f"/api/analytics/map?{Q}&metric=breaches&level=caza&condition_id={WARPLANE}").status_code == 422
    response = client.get(f"/api/analytics/map?{Q}&metric=breaches&level=village&condition_id={DRONE}")
    assert response.status_code == 422 and "village" in response.json()["detail"]
    assert client.get(f"/api/analytics/map?{Q}&metric=breaches&level=mohafaza&condition_id={WARPLANE}").status_code == 200


def test_map_air_values_follow_level_rules(client, data):
    moh = {v["code"]: v["value"] for v in client.get(f"/api/analytics/map?{Q}&metric=breaches&level=mohafaza").json()["values"]}
    assert moh["LB6"] == 3  # warplane + drone + helicopter (Sour); balloon sits in Aley / Mount Lebanon
    assert moh["LB3"] == 1
    caza = {v["code"]: v["value"] for v in client.get(f"/api/analytics/map?{Q}&metric=breaches&level=caza").json()["values"]}
    assert caza == {"LB63": 2, "LB31": 1}  # no warplane at caza level
    village = client.get(f"/api/analytics/map?{Q}&metric=breaches&level=village").json()["values"]
    assert sum(v["value"] for v in village) == 1


def test_kpis(client, data):
    body = client.get(f"/api/analytics/kpis?{Q}").json()
    assert body["incidents"]["current"] == 5
    assert body["deaths"]["current"] == 6 and body["deaths"]["male"] == 1 and body["deaths"]["unspecified"] == 3
    assert body["injuries"]["current"] == 6
    assert body["pending_review"] == 1
    assert body["unattributed_shared_toll_incidents"] == 1
    assert body["air_breaches"]["current"] == 4
    assert {r["condition_id"] for r in body["air_breaches"]["by_type"]} == {34, 35, 36, 38}
    assert body["villages_affected"]["current"] == 3 and body["cazas_affected"]["current"] == 2
    assert body["escalation"]["today"] == 1


def test_timeline_series_zero_filled_with_moving_average(client, data):
    body = client.get(f"/api/analytics/timeline?{Q}&grain=day").json()
    assert [p["date"] for p in body["total"]] == ["2026-10-01", "2026-10-02", "2026-10-03"]
    assert [p["value"] for p in body["total"]] == [2, 2, 1]
    assert body["coverage"]["available"] in (True, False)
    assert {s["key"] for s in body["series"]} >= {"airstrikes", "artillery", "gunfire"}
    month = client.get(f"/api/analytics/timeline?{Q}&grain=month").json()
    assert [p["value"] for p in month["total"]] == [5]
    assert client.get(f"/api/analytics/timeline?{Q}&grain=year").status_code == 422
    breaches = client.get(f"/api/analytics/timeline?{Q}&metric=breaches").json()
    assert sum(p["value"] for p in breaches["total"]) == 4


def test_rankings(client, data):
    body = client.get(f"/api/analytics/rankings?{Q}&level=village&limit=2").json()
    assert len(body["rows"]) == 2 and body["rows"][0]["value"] >= body["rows"][1]["value"]
    top = body["rows"][0]
    assert top["incidents"] == 3 and top["dominant_family"] == "airstrikes" and top["first_date"] == "2026-10-01"
    assert client.get(f"/api/analytics/rankings?{Q}&limit=500").status_code == 422
    caza = client.get(f"/api/analytics/rankings?{Q}&level=caza&metric=deaths").json()["rows"]
    assert caza[0]["code"] == "LB63" and caza[0]["value"] == 6
    air = client.get(f"/api/analytics/rankings?{Q}&level=mohafaza&metric=breaches").json()["rows"]
    assert air[0]["code"] == "LB6" and air[0]["value"] == 3


def test_peaks(client, data):
    body = client.get(f"/api/analytics/peaks?{Q}").json()
    assert body["incidents"]["day"]["value"] == 2
    assert body["deaths"]["day"]["start"] == "2026-10-01" and body["deaths"]["day"]["value"] == 5
    assert body["breaches"]["week"]["value"] == 4
    assert body["incidents"]["day"]["top_areas"][0]["code"] == "LB63"


def test_matrix(client, data):
    body = client.get(f"/api/analytics/matrix?{Q}").json()
    assert len(body["rows"]) == 26 and body["grand_total"] == 5
    sour = next(r for r in body["rows"] if r["code"] == "LB63")
    assert sour["cells"]["airstrikes"] == 3 and sour["cells"]["artillery"] == 1 and sour["total"] == 4
    assert body["column_totals"]["gunfire"] == 1


def test_area_profile(client, data):
    body = client.get(f"/api/analytics/area/LB63?{Q}").json()
    assert body["area"]["level"] == "caza" and body["totals"]["incidents"] == 4
    assert body["air_breaches"]["total"] == 2
    assert {f["code"] for f in body["families"]} == {"airstrikes", "artillery"}
    assert body["top_children"]["level"] == "village" and body["top_children"]["rows"][0]["incidents"] == 3
    assert body["first_last_hit"]["first_date"] == "2026-10-01"
    assert client.get("/api/analytics/area/NOPE").status_code == 404


def test_drilldown_matches_counts_and_paginates(client, data):
    body = client.get(f"/api/analytics/incidents?{Q}&page_size=100").json()
    assert body["total"] == 5 and len(body["items"]) == 5
    item = body["items"][0]
    assert set(item) >= {"id", "event_date", "village_en", "village_ar", "action_en", "verification_status",
                         "deaths", "injuries", "khabar_excerpt"}
    assert all(len(i["khabar_excerpt"]) <= 200 for i in body["items"])
    assert client.get(f"/api/analytics/incidents?{Q}&page_size=101").status_code == 422
    page2 = client.get(f"/api/analytics/incidents?{Q}&page_size=2&page=2").json()
    assert len(page2["items"]) == 2 and page2["total"] == 5
    area = client.get(f"/api/analytics/incidents?{Q}&area_code=LB61").json()
    assert area["total"] == 1
    deaths = client.get(f"/api/analytics/incidents?{Q}&metric=deaths").json()
    assert deaths["total"] == 3
    assert client.get(f"/api/analytics/incidents?{Q}&metric=breaches").status_code == 422
    withp = client.get(f"/api/analytics/incidents?{Q}&include_pending=true").json()
    assert withp["total"] == 6 and {i["bucket"] for i in withp["items"]} == {"counted", "pending"}


def test_export_workbook(client, data):
    from io import BytesIO

    from openpyxl import load_workbook

    response = client.get(f"/api/analytics/export.xlsx?{Q}")
    assert response.status_code == 200 and "spreadsheetml" in response.headers["content-type"]
    wb = load_workbook(BytesIO(response.content))
    assert wb.sheetnames == ["Summary", "By area", "By action", "Timeline", "Methodology"]
    methodology = [r[0].value for r in wb["Methodology"].iter_rows(min_row=2)]
    assert "Headline (counted)" in methodology and "Archive cut-off" in methodology


def test_cache_hit_and_version_bump(client, data, fake_cache, monkeypatch):
    store, versions = fake_cache
    first = client.get(f"/api/analytics/kpis?{Q}")
    assert len(store) == 1
    from app.analytics.repositories import analytics_repository as repo

    def boom(*a, **k):
        raise AssertionError("query ran despite cache hit")

    original = repo.incident_totals
    monkeypatch.setattr(repo, "incident_totals", boom)
    assert client.get(f"/api/analytics/kpis?{Q}").json() == first.json()
    monkeypatch.setattr(repo, "incident_totals", original)
    before = versions["v"]
    refresh_service.bump_cache_version()  # worker bumps the version -> old keys are never read again
    assert versions["v"] == before + 1
    client.get(f"/api/analytics/kpis?{Q}")
    assert len(store) == 2 and all(k.startswith("analytics:") for k in store)
